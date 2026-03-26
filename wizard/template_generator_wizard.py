import base64
import io
from odoo import models, fields, api, _
from odoo.exceptions import UserError


class TemplateGeneratorWizard(models.TransientModel):
    _name = 'template.generator.wizard'
    _description = 'Import Template Generator'

    model_id = fields.Many2one(
        'ir.model',
        string='Target Model',
        required=True,
        domain=[('transient', '=', False)],
        help='Generate an import template for this model.',
    )
    model_name = fields.Char(related='model_id.model')

    ruleset_id = fields.Many2one(
        'field.validation.ruleset',
        string='Rule Set (optional)',
        domain="[('model_id', '=', model_id)]",
        help='If specified, only fields with rules in this set are included.',
    )

    file_format = fields.Selection([
        ('xlsx', 'Excel (.xlsx) — recommended'),
        ('csv',  'CSV'),
    ], string='Output Format', default='xlsx', required=True)

    include_all_fields = fields.Boolean(
        string='Include All Model Fields',
        default=False,
        help='If unchecked, only fields that have at least one validation rule are included.',
    )
    include_example_row = fields.Boolean(
        string='Include Example Row',
        default=True,
    )
    include_rules_sheet = fields.Boolean(
        string='Include "Validation Rules" Sheet (Excel only)',
        default=True,
    )

    result_file = fields.Binary(string='Download Template', readonly=True, attachment=False)
    result_filename = fields.Char(string='Filename', readonly=True)
    state = fields.Selection([
        ('config',    'Configure'),
        ('download',  'Download'),
    ], default='config', readonly=True)

    @api.onchange('model_id')
    def _onchange_model_id(self):
        self.ruleset_id = False

    def action_generate(self):
        self.ensure_one()
        if self.file_format == 'xlsx':
            data, filename = self._generate_xlsx()
        else:
            data, filename = self._generate_csv()

        self.write({
            'result_file': base64.b64encode(data),
            'result_filename': filename,
            'state': 'download',
        })
        return self._reopen()

    def action_back(self):
        self.state = 'config'
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _generate_xlsx(self):
        try:
            import openpyxl
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side
            )
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.datavalidation import DataValidation
        except ImportError:
            raise UserError(_('openpyxl is required: pip install openpyxl'))

        Rule = self.env['field.validation.rule']
        rules_by_field = Rule.get_rules_for_model(self.model_name)

        if self.include_all_fields:
            model_fields = self.env['ir.model.fields'].search([
                ('model_id', '=', self.model_id.id),
                ('ttype', 'not in', ['one2many', 'many2many', 'binary', 'html']),
            ], order='field_description')
            columns = []
            for f in model_fields:
                col = {
                    'field_name': f.name,
                    'field_label': f.field_description,
                    'field_ttype': f.ttype,
                    'rules': rules_by_field.get(f.name, {}).get('rules', []),
                    'required': any(r['rule_type'] == 'required' for r in rules_by_field.get(f.name, {}).get('rules', [])),
                    'allowed_values': [],
                }
                for r in col['rules']:
                    if r['rule_type'] == 'allowed_values':
                        col['allowed_values'] = r['allowed_values']
                columns.append(col)
        else:
            if self.ruleset_id:
                field_names = list({r.field_name for r in self.ruleset_id.rule_ids})
            else:
                field_names = list(rules_by_field.keys())
            columns = []
            for fn in field_names:
                info = rules_by_field.get(fn, {})
                col = {
                    'field_name': fn,
                    'field_label': info.get('field_label', fn),
                    'field_ttype': info.get('field_ttype', 'char'),
                    'rules': info.get('rules', []),
                    'required': any(r['rule_type'] == 'required' for r in info.get('rules', [])),
                    'allowed_values': [],
                }
                for r in col['rules']:
                    if r['rule_type'] == 'allowed_values':
                        col['allowed_values'] = r['allowed_values']
                columns.append(col)

        wb = openpyxl.Workbook()

        ws = wb.active
        ws.title = 'Import Data'
        ws.sheet_view.showGridLines = True

        HEADER_REQUIRED = PatternFill('solid', fgColor='C0392B')
        HEADER_OPTIONAL = PatternFill('solid', fgColor='2980B9')
        HEADER_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        HINT_FILL = PatternFill('solid', fgColor='EBF5FB')
        HINT_FONT = Font(name='Calibri', italic=True, color='555555', size=9)
        EXAMPLE_FILL = PatternFill('solid', fgColor='EAFAF1')
        thin = Side(style='thin', color='BDBDBD')
        BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

        row_header = 1
        row_hint = 2
        row_example = 3 if self.include_example_row else None
        data_start_row = row_example + 1 if row_example else row_hint + 1

        ws.freeze_panes = f'A{data_start_row}'

        for col_idx, col in enumerate(columns, start=1):
            col_letter = get_column_letter(col_idx)

            hdr_cell = ws.cell(row=row_header, column=col_idx, value=col['field_label'])
            hdr_cell.fill = HEADER_REQUIRED if col['required'] else HEADER_OPTIONAL
            hdr_cell.font = HEADER_FONT
            hdr_cell.alignment = Alignment(horizontal='center', wrap_text=True)
            hdr_cell.border = BORDER

            hints = []
            for r in col['rules']:
                if r.get('hint_text'):
                    hints.append(r['hint_text'])
                elif r['rule_type'] == 'required':
                    hints.append('Required')
                elif r['rule_type'] == 'regex' and r.get('regex_pattern'):
                    hints.append(f'Pattern: {r["regex_pattern"]}')
                elif r['rule_type'] == 'allowed_values' and r.get('allowed_values'):
                    hints.append('Values: ' + ', '.join(r['allowed_values'][:4]))
                elif r['rule_type'] == 'min_value':
                    hints.append(f'Min: {r["param_min"]}')
                elif r['rule_type'] == 'max_value':
                    hints.append(f'Max: {r["param_max"]}')
                elif r['rule_type'] == 'email':
                    hints.append('Valid email address')
                elif r['rule_type'] == 'date_format' and r.get('date_format'):
                    hints.append(f'Format: {r["date_format"]}')

            hint_cell = ws.cell(row=row_hint, column=col_idx, value=' | '.join(hints) if hints else col['field_ttype'])
            hint_cell.fill = HINT_FILL
            hint_cell.font = HINT_FONT
            hint_cell.alignment = Alignment(horizontal='left', wrap_text=True)
            hint_cell.border = BORDER

            if self.include_example_row:
                example = self._get_example_value(col)
                ex_cell = ws.cell(row=row_example, column=col_idx, value=example)
                ex_cell.fill = EXAMPLE_FILL
                ex_cell.font = Font(name='Calibri', italic=True, color='2E7D32', size=10)
                ex_cell.border = BORDER

            if col['allowed_values'] and len(col['allowed_values']) <= 255:
                dv_formula = '"' + ','.join(col['allowed_values'][:20]) + '"'
                dv = DataValidation(
                    type='list',
                    formula1=dv_formula,
                    allow_blank=not col['required'],
                    showErrorMessage=True,
                    errorTitle='Invalid Value',
                    error=f'Please select a value from the list: {", ".join(col["allowed_values"][:5])}',
                    showDropDown=False,
                )
                dv.sqref = f'{col_letter}{data_start_row}:{col_letter}10000'
                ws.add_data_validation(dv)

            ws.column_dimensions[col_letter].width = max(18, min(40, len(col['field_label']) + 4))

        ws.row_dimensions[row_header].height = 28
        ws.row_dimensions[row_hint].height = 20

        ws.row_dimensions[data_start_row].height = 15

        legend_col = len(columns) + 2
        lcl = get_column_letter(legend_col)
        ws.cell(row=1, column=legend_col, value='LEGEND').font = Font(bold=True)
        ws.cell(row=2, column=legend_col, value='🔴 Required field').fill = PatternFill('solid', fgColor='FDECEA')
        ws.cell(row=3, column=legend_col, value='🔵 Optional field').fill = PatternFill('solid', fgColor='E8F4FD')
        ws.cell(row=4, column=legend_col, value='Row 2 = Hints').fill = HINT_FILL
        ws.cell(row=5, column=legend_col, value='Row 3 = Example (delete before import)').fill = EXAMPLE_FILL
        ws.column_dimensions[lcl].width = 38

        if self.include_rules_sheet:
            ws_rules = wb.create_sheet('Validation Rules')
            ws_rules.freeze_panes = 'A2'
            rule_headers = ['Field', 'Rule Name', 'Rule Type', 'Severity', 'Description / Error Message', 'Hint']
            for ci, h in enumerate(rule_headers, 1):
                c = ws_rules.cell(row=1, column=ci, value=h)
                c.fill = PatternFill('solid', fgColor='2C3E50')
                c.font = Font(bold=True, color='FFFFFF')
                c.alignment = Alignment(horizontal='center')

            row = 2
            for col in columns:
                for r in col['rules']:
                    ws_rules.cell(row=row, column=1, value=col['field_label'])
                    ws_rules.cell(row=row, column=2, value=r.get('rule_name', r['rule_type']))
                    ws_rules.cell(row=row, column=3, value=r['rule_type'])
                    sev = r['severity']
                    sev_cell = ws_rules.cell(row=row, column=4, value=sev.upper())
                    sev_cell.fill = PatternFill('solid', fgColor={
                        'error': 'FADBD8', 'warning': 'FDEBD0', 'info': 'D6EAF8'
                    }.get(sev, 'FFFFFF'))
                    ws_rules.cell(row=row, column=5, value=r['error_message'])
                    ws_rules.cell(row=row, column=6, value=r['hint_text'])
                    row += 1

            for ci, w in enumerate([20, 20, 18, 10, 60, 30], 1):
                ws_rules.column_dimensions[get_column_letter(ci)].width = w

        out = io.BytesIO()
        wb.save(out)
        model_label = self.model_id.name.replace(' ', '_').lower()
        filename = f'import_template_{model_label}.xlsx'
        return out.getvalue(), filename

    def _generate_csv(self):
        Rule = self.env['field.validation.rule']
        rules_by_field = Rule.get_rules_for_model(self.model_name)

        field_names = list(rules_by_field.keys())
        labels = [rules_by_field[fn]['field_label'] for fn in field_names]

        output = io.StringIO()
        import csv
        writer = csv.writer(output)
        writer.writerow(labels)
        hints = []
        for fn in field_names:
            parts = []
            for r in rules_by_field[fn]['rules']:
                if r['rule_type'] == 'required':
                    parts.append('REQUIRED')
                elif r.get('hint_text'):
                    parts.append(r['hint_text'])
            hints.append(' | '.join(parts))
        writer.writerow(['# HINTS: ' + h if h else '' for h in hints])

        if self.include_example_row:
            examples = []
            for fn in field_names:
                col = {
                    'field_name': fn,
                    'field_ttype': rules_by_field[fn].get('field_ttype', 'char'),
                    'rules': rules_by_field[fn]['rules'],
                    'allowed_values': next(
                        (r['allowed_values'] for r in rules_by_field[fn]['rules'] if r.get('allowed_values')), []
                    ),
                }
                examples.append(self._get_example_value(col))
            writer.writerow(examples)

        model_label = self.model_id.name.replace(' ', '_').lower()
        filename = f'import_template_{model_label}.csv'
        return output.getvalue().encode('utf-8-sig'), filename

    def _get_example_value(self, col):
        allowed = col.get('allowed_values', [])
        if allowed:
            return allowed[0]
        ttype = col.get('field_ttype', 'char')
        fn = col.get('field_name', '').lower()
        if 'email' in fn:
            return 'example@company.com'
        if 'phone' in fn or 'mobile' in fn:
            return '+1 555-0100'
        if 'date' in fn:
            return '2024-01-15'
        if ttype in ('integer', 'float', 'monetary'):
            return '0'
        if ttype == 'boolean':
            return 'True'
        if ttype == 'selection':
            return 'draft'
        return 'example_value'
